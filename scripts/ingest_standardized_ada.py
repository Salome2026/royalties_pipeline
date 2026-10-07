import hashlib
import re
import shutil
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import polars as pl


SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.append(str(SCRIPT_DIR))

from lib.statement_period import from_ada_filename
from lib.distributor_policy_store import load_distributor_policy_document
from lib.ada_identity import ada_isrc_expr, ada_native_code_expr, add_ada_artist_evidence
from lib.ada_release_codes import add_ada_release_codes


BASE = Path(r"C:\royalties_pipeline")
INPUT_ROOT = BASE / "input_raw" / "ada"
OUTPUT_PATH = BASE / "warehouse" / "marts" / "standardized_raw_ada.parquet"
TEMP_DIR = BASE / "staging" / "standardized_raw_parts" / "ada"

SOURCE = "ada"
STATEMENT_TYPE = "ada_royalty_detail"
SOURCE_SHEET = "royalty_detail"

ACCOUNTS = {
    "mawz": {
        "input_directory": "mawz",
        "original_account": "99205",
    },
    "indyana_records": {
        "input_directory": "Indyana Records",
        "original_account": "99500",
    },
}


def load_ada_policy_rules() -> dict[tuple[str, str], dict]:
    payload = load_distributor_policy_document()
    rules: dict[tuple[str, str], dict] = {}
    for entry in payload.get("entries", []):
        if str(entry.get("source") or "").strip().lower() != SOURCE:
            continue
        account = str(entry.get("account") or "").strip().lower()
        for source_sheet, rule in (entry.get("sheet_rules") or {}).items():
            if isinstance(rule, dict):
                rules[(account, str(source_sheet or "").strip())] = rule
    return rules


def policy_bool(value) -> bool:
    return value is True


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def text_expr(name: str, columns: set[str]) -> pl.Expr:
    if name not in columns:
        return pl.lit(None).cast(pl.Utf8)
    return pl.col(name).cast(pl.Utf8, strict=False).str.strip_chars().replace("", None)


def decimal_expr(name: str, columns: set[str]) -> pl.Expr:
    if name not in columns:
        return pl.lit(None).cast(pl.Float64)
    return (
        pl.col(name)
        .cast(pl.Utf8, strict=False)
        .str.replace_all(",", "")
        .str.strip_chars()
        .replace("", None)
        .cast(pl.Float64, strict=False)
    )


def read_statement(path: Path) -> pl.DataFrame:
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"ADA admite solamente Excel .xlsx: {path.name}")
    return read_excel_statement(path)


def read_excel_statement(path: Path) -> pl.DataFrame:
    from openpyxl import load_workbook

    period = from_ada_filename(path.name).period
    match = re.fullmatch(r"(\d+)_\d{6}_\d{6}_\1_DTL\.xlsx", path.name, flags=re.IGNORECASE)
    if period == "unknown" or match is None:
        raise ValueError(f"Nombre de statement Excel ADA no valido: {path.name}")
    account = match.group(1)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Distribution Statement" not in workbook.sheetnames:
            raise ValueError(f"Falta la hoja Distribution Statement: {path.name}")
        sheet = workbook["Distribution Statement"]
        rows = iter(sheet.iter_rows(values_only=True))
        required = {"Artist Name", "Project Title", "Catalogue Number", "ISRC", "Catalogue Title", "Reported Month", "Territory", "Country Code", "Price Name", "Revenue Type Desc", "Sales", "Receipts Value", "Distribution Fees", "Artist Royalties", "Mechanical Fees", "Admin Fees", "Upload Fees", "Other Fees"}
        contract = None
        workbook_period = None
        headers = None
        for row_number, row in enumerate(rows, start=1):
            names = [str(value).strip() if value is not None else "" for value in row]
            for index, name in enumerate(names[:-1]):
                if name.rstrip(":") == "Contract":
                    contract = names[index + 1]
                if name.rstrip(":") == "Royalties Statement for Period":
                    workbook_period = names[index + 1]
            if required.issubset(set(names)):
                headers = names
                break
            if row_number >= 20:
                break
        if headers is None:
            raise ValueError(f"No se encontro la cabecera de detalle ADA: {path.name}")
        if not contract or not re.match(rf"^{re.escape(account)}\s*-", contract):
            raise ValueError(f"Contrato del Excel no coincide con la cuenta {account}: {path.name}")
        expected_period = datetime.strptime(period, "%Y-%m").strftime("%m/%y")
        if not workbook_period or re.sub(r"\s+", "", workbook_period) != f"{expected_period}-{expected_period}":
            raise ValueError(f"Periodo del Excel no coincide con {period}: {path.name}")
        positions = [(index, name) for index, name in enumerate(headers) if name]
        if len({name for _, name in positions}) != len(positions):
            raise ValueError(f"Columnas duplicadas en ADA: {path.name}")
        records = []
        footer_labels = {"sub totals", "previously accounted deductions", "totals"}
        for row in rows:
            nonempty = [value for value in row if value is not None and str(value).strip()]
            if not nonempty:
                continue
            first_label = str(nonempty[0]).strip().rstrip(":").lower()
            has_identity = any(row[index] is not None and str(row[index]).strip()
                               for index, name in positions if name in {"ISRC", "Catalogue Number"})
            if first_label in footer_labels and not has_identity:
                continue
            records.append([str(row[index]).strip() if row[index] is not None else None for index, _ in positions])
        frame = pl.DataFrame(records, schema={name: pl.Utf8 for _, name in positions}, orient="row")
    finally:
        workbook.close()

    columns = set(frame.columns)
    # Reject unknown financial components instead of silently losing money.
    fee_names = ["Distribution Fees", "Mechanical Fees", "Admin Fees", "Upload Fees", "Other Fees"]
    for name in ["Receipts Value", "Sales", "Artist Royalties", *fee_names]:
        value = decimal_expr(name, columns)
        if frame.select((value.is_null() | ~value.is_finite()).any()).item():
            raise ValueError(f"{name} contiene valores vacios o invalidos: {path.name}")
        if name in ["Artist Royalties", *fee_names[1:]] and frame.select((value != 0).any()).item():
            raise ValueError(f"{name} no es cero; revisar la regla economica de ADA: {path.name}")
    gross = decimal_expr("Receipts Value", columns)
    exact_fees = [Decimal(value.replace(",", "")) for value in frame.get_column("Distribution Fees")]
    exact_net = [float(Decimal(value.replace(",", "")) - fee) for value, fee in zip(frame.get_column("Receipts Value"), exact_fees)]
    consumed = text_expr("Reported Month", columns).str.replace(r"^(\d{4})(\d{2})$", "${1}-${2}")
    frame = frame.with_columns([
        text_expr("Catalogue Number", columns).alias("catalog_number"),
        text_expr("Parent Product ID", columns).alias("parent_product_id"),
        text_expr("UPC", columns).alias("ada_explicit_upc"),
        pl.lit(account).alias("ada_account_id"),
        pl.lit(period).alias("statement_period"),
        consumed.alias("transaction_month"),
        gross.alias("gross_royalty_usd"),
        pl.Series("deductible_fees_usd", [float(value) for value in exact_fees], dtype=pl.Float64),
        pl.Series("amount_usd", exact_net, dtype=pl.Float64),
    ])
    months = frame.get_column("transaction_month").unique().to_list()
    for month in months:
        if month is None:
            raise ValueError(f"Mes de consumo vacio en ADA: {path.name}")
        datetime.strptime(month, "%Y-%m")
    return frame


def select_statement_files(input_dir: Path) -> list[Path]:
    statements: dict[str, Path] = {}
    for path in sorted(input_dir.iterdir()):
        suffix = path.suffix.lower()
        if not path.is_file() or path.name.startswith("~$"):
            continue
        if suffix != ".xlsx":
            raise ValueError(f"ADA admite solamente Excel .xlsx; retirar de ingesta {path.name}")
        period = from_ada_filename(path.name).period
        if period == "unknown":
            raise ValueError(f"No se pudo identificar el statement ADA: {path.name}")
        if period in statements:
            raise ValueError(f"Statements ADA duplicados para {period}: {statements[period].name}, {path.name}")
        statements[period] = path
    return [path for _, path in sorted(statements.items())]


def standardize(
    frame: pl.DataFrame,
    path: Path,
    account: str,
    expected_original_account: str,
    policy_rule: dict,
) -> pl.DataFrame:
    frame = add_ada_release_codes(frame)
    columns = set(frame.columns)
    required = {
        "transaction_month",
        "ISRC",
        "Catalogue Title",
        "Artist Name",
        "Territory",
        "Sales",
        "gross_royalty_usd",
        "deductible_fees_usd",
        "amount_usd",
    }
    missing = sorted(required - columns)
    if missing:
        raise ValueError(f"Faltan columnas requeridas ADA: {missing}")
    if frame.select((ada_isrc_expr(columns).is_null() & ada_native_code_expr(columns).is_null()).any()).item():
        raise ValueError(f"ADA contiene ingresos sin ISRC ni identificador nativo: {path.name}")

    original_accounts = {
        str(value).strip()
        for value in frame.get_column("ada_account_id").drop_nulls().unique().to_list()
    }
    if frame.height and original_accounts != {expected_original_account}:
        raise ValueError(
            f"La carpeta ADA/{account} esperaba Account={expected_original_account}, "
            f"pero {path.name} contiene {sorted(original_accounts)}"
        )

    statement = from_ada_filename(path.name)
    if statement.period == "unknown":
        raise ValueError(f"No se pudo obtener statement_period de {path.name}")

    for period_column in ["statement_period"]:
        periods = (
            frame.select(text_expr(period_column, columns).drop_nulls().unique())
            .to_series()
            .to_list()
        )
        unexpected = sorted({str(value) for value in periods if str(value) != statement.period})
        if unexpected:
            raise ValueError(
                f"{period_column} no coincide con filename {statement.period}: {unexpected}"
            )

    file_hash = sha256_file(path)
    ingested_at = datetime.now().isoformat(timespec="seconds")
    gross = decimal_expr("gross_royalty_usd", columns)
    fees = decimal_expr("deductible_fees_usd", columns)
    net = decimal_expr("amount_usd", columns)
    valid_amounts = pl.all_horizontal([value.is_not_null() & value.is_finite() for value in [gross, fees, net]])
    if frame.select((~valid_amounts | ((gross - fees - net).abs() > 1e-8)).any()).item():
        raise ValueError(f"ADA contiene importes invalidos o bruto menos fees distinto del neto: {path.name}")
    consumption = text_expr("transaction_month", columns)
    if frame.select((consumption.is_null() | ~consumption.str.contains(r"^[0-9]{4}-(0[1-9]|1[0-2])$")).any()).item():
        raise ValueError(f"ADA contiene un mes de consumo invalido: {path.name}")
    revenue_basis = str(policy_rule.get("revenue_basis") or "").strip()
    if not revenue_basis:
        raise ValueError(f"La policy ADA/{account}/{SOURCE_SHEET} no define revenue_basis.")

    return frame.with_columns([
        net.alias("amount_usd"),
        net.alias("net_amount"),
        net.alias("net_amount_usd"),
        gross.alias("gross_royalty_usd"),
        fees.alias("deductible_fees_usd"),
        pl.lit("USD").alias("currency_original"),

        consumption.alias("transaction_month"),
        pl.lit(None).cast(pl.Utf8).alias("receipt_month"),
        pl.lit(statement.period).alias("statement_period"),
        pl.lit(statement.source).alias("statement_period_source"),
        pl.lit(statement.note).alias("statement_period_note"),

        text_expr("Artist Name", columns).alias("artist_statement_style"),
        text_expr("Catalogue Title", columns).alias("track_statement_style"),
        text_expr("Catalogue Title", columns).alias("asset_title_statement"),
        text_expr("Project Title", columns).alias("release_statement_style"),
        ada_isrc_expr(columns).alias("asset_isrc"),
        text_expr("product_upc", columns).alias("product_upc"),
        text_expr("Catalogue Number", columns).alias("catalog_number"),
        ada_native_code_expr(columns).alias("source_asset_id"),
        pl.lit(expected_original_account).alias("ada_account_id"),
        text_expr("parent_product_id", columns).alias("parent_product_id"),
        pl.lit("excel").alias("ada_statement_format"),
        text_expr("Local Product Number", columns).alias("local_product_number"),

        text_expr("Territory", columns).alias("store_name"),
        text_expr("Country Code", columns).alias("territory"),
        text_expr("Revenue Type Desc", columns).alias("use_type"),
        text_expr("Price Name", columns).alias("product_type"),
        decimal_expr("Sales", columns).alias("units"),

        pl.lit(SOURCE).alias("source"),
        pl.lit(account).alias("account"),
        pl.lit(STATEMENT_TYPE).alias("statement_type"),
        pl.lit(SOURCE_SHEET).alias("source_sheet"),
        pl.lit(revenue_basis).alias("revenue_basis"),
        pl.lit(policy_bool(policy_rule.get("cash_view"))).alias("include_in_cash_view"),
        pl.lit(policy_bool(policy_rule.get("catalog_view"))).alias("include_in_catalog_view"),
        pl.lit(policy_bool(policy_rule.get("statement_view"))).alias("include_in_statement_view"),
        pl.lit(revenue_basis == "transfer").alias("possible_internal_transfer"),

        pl.lit(path.name).alias("statement_file_name"),
        pl.lit(str(path)).alias("statement_file_path"),
        pl.lit(file_hash).alias("statement_file_hash"),
        pl.lit(ingested_at).alias("ingested_at"),
    ])


def main() -> None:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    if TEMP_DIR.exists():
        shutil.rmtree(TEMP_DIR)
    TEMP_DIR.mkdir(parents=True, exist_ok=True)

    policy_rules = load_ada_policy_rules()

    parts: list[Path] = []
    empty_periods: list[str] = []
    statement_count = 0

    for account, config in ACCOUNTS.items():
        input_dir = INPUT_ROOT / str(config["input_directory"])
        files = select_statement_files(input_dir)
        if not files:
            raise FileNotFoundError(f"No hay statements ADA/{account} en {input_dir}")

        policy_rule = policy_rules.get((account, SOURCE_SHEET))
        if policy_rule is None:
            raise ValueError(
                f"Falta policy ADA para account={account!r}, source_sheet={SOURCE_SHEET!r}. "
                "Actualizar la politica de distribuidoras en Cloud SQL antes de ingerir."
            )

        account_temp_dir = TEMP_DIR / account
        account_temp_dir.mkdir(parents=True, exist_ok=True)
        statement_count += len(files)
        print(f"ADA/{account}: {len(files)} statements")

        for path in files:
            statement = from_ada_filename(path.name)
            frame = read_statement(path)
            if frame.is_empty():
                empty_periods.append(f"{account}/{statement.period}")
                print(f"  {statement.period}: sin actividad")
                continue

            standardized = standardize(
                frame,
                path,
                account,
                str(config["original_account"]),
                policy_rule,
            )
            part_path = account_temp_dir / f"{path.stem}.parquet"
            standardized.write_parquet(part_path)
            parts.append(part_path)
            print(
                f"  {statement.period}: {standardized.height} filas | "
                f"USD {standardized['amount_usd'].sum():.8f}"
            )

    if not parts:
        raise ValueError("Todos los statements ADA estan vacios; no se genero el mart.")

    final = add_ada_artist_evidence(pl.concat([pl.read_parquet(path) for path in parts], how="diagonal_relaxed"))
    temporary_output = OUTPUT_PATH.with_name(f"{OUTPUT_PATH.name}.tmp")
    final.write_parquet(temporary_output)
    temporary_output.replace(OUTPUT_PATH)

    print(f"Archivo: {OUTPUT_PATH}")
    print(f"Cuentas: {len(ACCOUNTS)}")
    print(f"Statements revisados: {statement_count}")
    print(f"Statements con movimientos: {len(parts)}")
    print(f"Statements sin actividad: {len(empty_periods)} ({', '.join(empty_periods)})")
    print(f"Filas: {final.height}")
    print(f"Total bruto USD: {final['gross_royalty_usd'].sum()}")
    print(f"Fees USD: {final['deductible_fees_usd'].sum()}")
    print(f"Total neto USD: {final['amount_usd'].sum()}")


if __name__ == "__main__":
    main()
