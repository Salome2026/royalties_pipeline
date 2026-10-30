"""Report clarity across sources without changing compact grouping or amounts."""
from pathlib import Path
from types import SimpleNamespace
import sys
import unittest

import openpyxl
import pandas as pd
import polars as pl

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT), str(ROOT / "scripts")]

from scripts.build_keyword_royalty_report import build_track_summary_lf, display_dataframe, style_workbook
from scripts.lib.report_presentation import YOUTUBE_CHANNEL_LABEL, youtube_video_url
from app.royalty_reports.google_sheets import column_width, dataframe_values, video_code_link_request


def movement(**changes):
    return {"source": "fuga", "account": "test", "asset_isrc": "ARDL12500066",
            "video_id": "m6mrKKUuosk", "catalog_key": "ISRC:ARDL12500066",
            "report_code": "ARDL12500066", "report_code_source": "ISRC",
            "track_statement_style": "Tema", "artist_statement_style": "Artista",
            "dsp_normalized": "YouTube", "content_origin_normalized": "Video / Channel",
            "content_type": None, "amount_usd": 100., "units": 50., "period_month": "2026-01",
            **changes}


def summarize(rows):
    return build_track_summary_lf(pl.DataFrame(rows).lazy(), include_identity=True).collect()


class ReportVideoPresentationTests(unittest.TestCase):
    def test_resolved_evidence_works_without_source_specific_branches(self):
        # The input is an existing resolved dimension, not a new source inference.
        for source in ["fuga", "onerpm", "orchard", "dashgo", "ada", "soundon", "wml"]:
            with self.subTest(source=source):
                result = summarize([
                    movement(source=source, asset_isrc=None, catalog_key="VIDEO:m6mrKKUuosk"),
                    movement(source=source, period_month="2026-02", amount_usd=-10., units=-5.),
                    movement(source=source, content_origin_normalized="Music / Art Track", amount_usd=40.),
                    movement(source=source, content_origin_normalized="UGC / Content ID", amount_usd=20.),
                ])
                self.assertEqual(result.height, 2)
                video = result.filter(pl.col("report_code_source") == "Video ID").row(0, named=True)
                self.assertEqual((video["amount_usd"], video["units"]), (90., 45.))
                self.assertEqual((video["first_month"], video["last_month"]), ("2026-01", "2026-02"))
                self.assertEqual(video["content_types"], YOUTUBE_CHANNEL_LABEL)
                audio = result.filter(pl.col("report_code_source") == "ISRC").row(0, named=True)
                self.assertEqual((audio["amount_usd"], audio["units"]), (60., 100.))
                self.assertEqual(audio["content_types"], "")
                self.assertEqual(result["amount_usd"].sum(), 150.)

    def test_native_types_are_retained_and_unknowns_are_not_inferred_from_title(self):
        result = summarize([movement(content_type="catalog"), movement(content_type="ugc")])
        self.assertEqual(result["content_types"][0], YOUTUBE_CHANNEL_LABEL + " | catalog | ugc")
        for source in ["dashgo", "ada", "soundon", "orchard", "wml"]:
            with self.subTest(source=source):
                result = summarize([movement(source=source, track_statement_style="Tema (Video Oficial)",
                    content_origin_normalized="No informado", content_type="Internet Video")])
                self.assertEqual(result["content_types"][0], "Internet Video")

    def test_mixed_video_identity_is_not_presented_as_all_channel_income(self):
        rows = [movement(catalog_key="VIDEO:m6mrKKUuosk"),
                movement(catalog_key="VIDEO:m6mrKKUuosk", content_origin_normalized="No informado")]
        result = summarize(rows)
        self.assertEqual(result.height, 1)
        self.assertEqual(result["content_types"][0], "")
        self.assertEqual(result["amount_usd"][0], 200.)

    def test_missing_dimensions_keep_the_previous_output(self):
        row = movement(content_type="catalog")
        del row["dsp_normalized"]
        del row["content_origin_normalized"]
        result = summarize([row])
        self.assertEqual(result["content_types"][0], "catalog")
        self.assertEqual(result["report_code_source"][0], "ISRC")

    def test_only_explicit_valid_video_codes_become_links(self):
        self.assertEqual(youtube_video_url("m6mrKKUuosk", "Video ID"),
                         "https://www.youtube.com/watch?v=m6mrKKUuosk")
        for code, kind in [("m6mrKKUuosk", "ISRC"), ("m6mrKKUuosk", "ID"),
                           ("deleted-x", "Video ID"), ("m6mrKKUuosk|asset", "Video ID"),
                           (None, "Video ID"), ("m6mrKKUuosk", pd.NA)]:
            self.assertIsNone(youtube_video_url(code, kind))

    def test_excel_and_sheets_links_preserve_cells_columns_and_formats(self):
        summary = summarize([movement(), movement(content_origin_normalized="Music / Art Track")])
        table = display_dataframe(summary.drop("_report_identity").to_pandas())
        before = dataframe_values(table)
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.title = "resumen por tema"
        for row in before:
            ws.append(row)
        style_workbook(SimpleNamespace(book=workbook))
        self.assertEqual(list(ws.values), [tuple(row) for row in before])
        self.assertEqual((ws.max_row, ws.max_column), (3, 10))
        self.assertEqual(ws.auto_filter.ref, "A1:J3")
        self.assertEqual(ws.freeze_panes, "A2")
        headers = {c.value: c.column for c in ws[1]}
        video_index = next(i for i, row in enumerate(before[1:], 2) if row[3] == "Video ID")
        cell = ws.cell(video_index, headers["Codigo"])
        self.assertEqual(cell.value, "m6mrKKUuosk")
        self.assertEqual(cell.hyperlink.target, "https://www.youtube.com/watch?v=m6mrKKUuosk")
        self.assertEqual(ws.cell(video_index, headers["Ingresos USD"]).number_format, '$#,##0.00')
        request = video_code_link_request(table, 17)["updateCells"]
        self.assertEqual(request["start"], {"sheetId": 17, "rowIndex": 1, "columnIndex": 2})
        self.assertEqual(request["fields"], "userEnteredFormat.textFormat.link")
        self.assertEqual(request["rows"][video_index - 2]["values"][0]
                         ["userEnteredFormat"]["textFormat"]["link"]["uri"], cell.hyperlink.target)
        self.assertEqual(dataframe_values(table), before)
        self.assertGreaterEqual(column_width("Tipo de contenido", table), len(YOUTUBE_CHANNEL_LABEL) * 7 + 24)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
