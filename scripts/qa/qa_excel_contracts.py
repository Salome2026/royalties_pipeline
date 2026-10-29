"""Opt-in Excel enrichment, broad search preservation and legacy compatibility."""
from __future__ import annotations

import ast
from copy import deepcopy
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT), str(ROOT / "scripts"), str(ROOT / "scripts/qa")]

import openpyxl
import polars as pl
from fastapi import HTTPException
from pydantic import BaseModel, Field
from typing import Any, Literal
from app.royalty_reports.contracts import ReportRequest
from app.royalty_reports.excel_contracts import participant_columns, reference_fragments, require_snapshot
from app.master_contract_associations import association_key
from lib.distributor_policy_store import use_distributor_policy_snapshot
from qa_contractual_executive import fixture
import build_keyword_royalty_report as report

ISRC = "ARDL12600001"


def snapshot():
    previous, _, _, _ = fixture()
    return {"schema_version": 1, "kind": "excel_contract_columns", "release_id": "release-a", "as_of": "2026-10-29",
        "contracts": previous["contracts"], "catalog": [{"catalog_key": "ISRC:" + ISRC, "asset_isrc": ISRC,
            "catalog_artists": ["Proyecto", "Artista"], "artist_statement": "Proyecto and Artista"}],
        "associations": {"included": {association_key("VIDEO", "video-code1", "fuga", "indy"): [ISRC]}, "excluded": {}}}


def row(**changes):
    return {"source": "fuga", "account": "indy", "source_sheet": "royalty_detail", "revenue_basis": "generation",
        "statement_period": "2026-06", "transaction_month": "2026-05", "asset_isrc": ISRC,
        "product_upc": None, "video_id": None, "track_id": None, "catalog_number": None,
        "track_statement_style": "Tema", "artist_statement_style": "Proyecto and Artista", "amount_usd": 100.,
        "units": 10., "territory": "AR", "store": "spotify", "store_name": "Spotify", "usage_type": "streaming",
        "currency_original": "USD", "statement_file_name": "prueba.csv", **changes}


def evaluate(rows=None, state=None, scope="all", keywords=None):
    frame = report.add_report_code(pl.DataFrame(rows or [row()]).lazy(), set(pl.DataFrame(rows or [row()]).columns)).with_columns(pl.col("transaction_month").alias("period_month"))
    summary = report.build_track_summary_lf(frame, include_identity=True).collect()
    names, values = participant_columns(summary, reference_fragments(frame), state or snapshot(), keywords or ["Artista"], scope)
    return names, values, summary


def canonical_functions():
    import re
    tree = ast.parse((ROOT / "app/vpo_corp_api.py").read_text(encoding="utf-8"))
    selected = [node for node in tree.body if isinstance(node, (ast.ClassDef, ast.FunctionDef))
                and node.name in {"RoyaltyReportJobRequest", "canonical_royalty_report_params"}]
    namespace = {"BaseModel": BaseModel, "Field": Field, "Literal": Literal, "Any": Any, "re": re,
        "HTTPException": HTTPException, "normalize_keywords": lambda values: values, "is_wml": lambda: False}
    exec(compile(ast.Module(body=selected, type_ignores=[]), "excel-request-validation", "exec"), namespace)
    namespace["RoyaltyReportJobRequest"].model_rebuild(_types_namespace=namespace)
    return namespace


class ExcelContractTests(unittest.TestCase):
    def test_final_recipients_and_unpaid_credited_project(self):
        names, values, _ = evaluate()
        self.assertEqual(values[0]["Indyana"], (.54, 54.))
        self.assertEqual(values[0]["Socio"], (.4, 40.))
        self.assertEqual(values[0]["Artista"], (.06, 6.))
        self.assertEqual(values[0]["Proyecto"], (0., 0.))
        self.assertEqual(len(names), len(set(names)))

    def test_matching_selects_columns_not_income_or_rows(self):
        names, values, summary = evaluate(scope="matching", keywords=["arti", "nombre de tema"])
        self.assertEqual(names, ["Artista"])
        self.assertEqual(values[0]["Artista"], (.06, 6.))
        self.assertEqual(summary["amount_usd"].sum(), 100.)
        self.assertEqual(evaluate(scope="matching", keywords=[ISRC])[0], [])

    def test_associated_video_same_contract_but_separate_income(self):
        _, values, summary = evaluate([row(), row(asset_isrc=None, video_id="video-code1", amount_usd=50., track_statement_style="Video")])
        self.assertEqual(summary.height, 2)
        self.assertEqual(summary["amount_usd"].sum(), 150.)
        self.assertEqual(sorted(v["Artista"][1] for v in values), [3., 6.])

    def test_unconfirmed_and_excluded_association_are_blank(self):
        for excluded in (False, True):
            state = snapshot()
            key = association_key("VIDEO", "video-code1", "fuga", "indy")
            if excluded:
                state["associations"]["excluded"][key] = [ISRC]
            else:
                state["associations"]["included"] = {}
            names, values, summary = evaluate([row(asset_isrc=None, video_id="video-code1")], state)
            self.assertIn("Artista", names)
            self.assertTrue(all(value is None for value in values[0].values()))
            self.assertEqual(summary.height, 1)

    def test_ada_native_product_is_not_allocated_as_track(self):
        state = snapshot()
        state["associations"]["included"][association_key("UPC", "123456789012", "ada", "indy")] = [ISRC]
        _, values, _ = evaluate([row(source="ada", asset_isrc=None, catalog_number="ADA-ALBUM", product_upc="123456789012")], state)
        self.assertTrue(all(value is None for value in values[0].values()))

    def test_open_absent_and_invalid_contracts_leave_blanks_not_zero(self):
        for kind in ("open", "absent", "invalid"):
            state = snapshot()
            if kind == "open": state["contracts"][0]["is_closed"] = False
            elif kind == "absent": state["contracts"] = []
            else: state["contracts"][0]["split"]["agreements"][0]["master_pool_percent"] = 90
            _, values, summary = evaluate(state=state)
            self.assertTrue(all(value is None for value in values[0].values()), kind)
            self.assertEqual(summary["amount_usd"].sum(), 100.)

    def test_missing_statement_and_outside_contract_are_blank(self):
        for month in (None, "2025-12"):
            self.assertTrue(all(value is None for value in evaluate([row(statement_period=month)])[1][0].values()))

    def test_partial_coverage_uses_effective_share_without_extra_columns(self):
        _, values, summary = evaluate([row(statement_period="2025-12"), row(statement_period="2026-01")])
        self.assertEqual(summary["amount_usd"].sum(), 200.)
        self.assertEqual(values[0]["Artista"], (.03, 6.))

    def test_contiguous_contracts_use_statement_not_consumption(self):
        state = snapshot()
        agreement = state["contracts"][0]["split"]["agreements"][0]
        agreement["effective_until"] = "2026-01-31"
        second = deepcopy(agreement)
        second.update(id="b", effective_from="2026-02-01", effective_until=None)
        second["allocation"]["participants"][0]["rule"]["retained_percent"] = 50
        state["contracts"][0]["split"]["agreements"].append(second)
        _, values, _ = evaluate([row(statement_period="2026-01", transaction_month="2025-12"),
                                row(statement_period="2026-02", transaction_month="2025-12")], state)
        self.assertEqual(values[0]["Artista"], (.08, 16.))

    def test_refunds_and_cents_conserve_eligible_income(self):
        for amount in (-100., .01, -.01, 100.005, 0.):
            _, values, _ = evaluate([row(amount_usd=amount)])
            self.assertAlmostEqual(sum(pair[1] for pair in values[0].values() if pair is not None), amount)

    def test_zero_net_with_changed_contract_does_not_invent_effective_percentage(self):
        state = snapshot()
        agreement = state["contracts"][0]["split"]["agreements"][0]
        agreement["effective_until"] = "2026-01-31"
        second = deepcopy(agreement)
        second.update(id="b", effective_from="2026-02-01", effective_until=None)
        second["allocation"]["participants"][0]["rule"]["retained_percent"] = 50
        state["contracts"][0]["split"]["agreements"].append(second)
        _, values, _ = evaluate([row(statement_period="2026-01", amount_usd=100.), row(statement_period="2026-02", amount_usd=-100.)], state)
        self.assertEqual(values[0]["Artista"], (None, -4.))
        self.assertEqual(sum(value[1] for value in values[0].values()), 0.)

    def test_reconciliation_difference_preserves_income_but_does_not_invent_split(self):
        frame = pl.DataFrame([row()]).lazy()
        frame = report.add_report_code(frame, set(frame.collect_schema().names())).with_columns(pl.col("transaction_month").alias("period_month"))
        summary = report.build_track_summary_lf(frame, include_identity=True).collect().with_columns(pl.lit(101.).alias("amount_usd"))
        _, values = participant_columns(summary, reference_fragments(frame), snapshot(), ["Artista"], "all")
        self.assertTrue(all(value is None for value in values[0].values()))

    def test_snapshot_required_no_live_contract_fallback(self):
        for value in (None, {}, {"schema_version": 2}, {**snapshot(), "associations": {}}):
            with self.assertRaises(ValueError): require_snapshot(value)

    def test_legacy_request_canonical_parameters_exactly_preserved(self):
        ns = canonical_functions()
        model, canonical = ns["RoyaltyReportJobRequest"], ns["canonical_royalty_report_params"]
        old = canonical(model(keywords=["Artista"]), {})
        self.assertNotIn("include_contracts", old)
        self.assertNotIn("contract_columns", old)
        self.assertEqual(old, canonical(model(keywords=["Artista"], include_contracts=False, contract_columns="matching"), {}))
        new = canonical(model(keywords=["ar"], include_contracts=True), {})
        self.assertTrue(new["include_contracts"])
        self.assertEqual(new["keywords"], ["ar"])
        self.assertNotIn("contract_artists", new)
        self.assertEqual(new["period_basis"], "transaction_month")

    def test_non_excel_and_wml_cannot_enable_enrichment(self):
        ns = canonical_functions()
        model, canonical = ns["RoyaltyReportJobRequest"], ns["canonical_royalty_report_params"]
        for output in ("google_sheet", "executive_pdf"):
            with self.assertRaises(HTTPException): canonical(model(output=output, keywords=["Artista"], include_contracts=True), {})
        ns["is_wml"] = lambda: True
        with self.assertRaises(HTTPException): canonical(model(keywords=["Artista"], include_contracts=True), {})

    def test_request_defaults_preserve_ordinary_jobs(self):
        job = {"id": 1, "report_key": "royalty_keyword", "output_format": "excel", "params": {"keywords": ["art"]}}
        self.assertFalse(ReportRequest.from_job(job, keywords=["art"]).include_contracts)
        job["params"].update(include_contracts=True, contract_columns="matching")
        self.assertEqual(ReportRequest.from_job(job, keywords=["art"]).contract_columns, "matching")

    def test_participant_name_cannot_overwrite_an_original_column(self):
        from app.royalty_reports.excel_contracts import enrich_tables
        frame = report.add_report_code(pl.DataFrame([row()]).lazy(), set(row())).with_columns(pl.col("transaction_month").alias("period_month"))
        summary = report.build_track_summary_lf(frame, include_identity=True).collect()
        table = report.display_dataframe(summary.drop("_report_identity").to_pandas())
        old = table.copy(deep=True)
        tables = {"resumen por tema": table}
        enrich_tables(tables, summary, frame, snapshot(), ["Artista"], "all")
        self.assertTrue(tables["resumen por tema"].iloc[:, :len(old.columns)].equals(old))
        self.assertEqual(list(tables["resumen por tema"].columns).count("Artista"), 2)

    def test_excel_changes_only_append_participant_pairs_and_styles(self):
        state = json.loads(json.dumps(snapshot()).replace('"Artista"', '"Invitada"'))
        policy = {"schema_version": 1, "policy_version": 1, "entries": [{"source": "fuga", "account": "indy",
            "sheet_rules": {"royalty_detail": {"catalog_view": True, "statement_view": True, "revenue_basis": "generation"}}}]}
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            song, detail = root / "song.parquet", root / "raw.parquet"
            pl.DataFrame([row(artist_statement_style="Proyecto and Invitada"), row(asset_isrc="ARDL12600002", artist_statement_style="Proyecto and Invitada", track_statement_style="Sin contrato", amount_usd=20.)]).write_parquet(song)
            detail.write_bytes(song.read_bytes())
            with patch.dict(os.environ, {"VPO_CATALOG_MASTER_PATH": str(root / "missing.parquet"), "VPO_CATALOG_STATUS_PATH": str(root / "missing-status.parquet")}), use_distributor_policy_snapshot(policy):
                for basis in ("statement_period", "transaction_month"):
                    args = dict(keywords=["Invitada"], mode="any", raw_limit=5000, song_path=song, standardized_path=detail, period_basis=basis)
                    original = report.build_report(**args, output_dir=root / (basis + "-old"))
                    enriched = report.build_report(**args, output_dir=root / (basis + "-new"), contract_snapshot=state)
                    a, b = openpyxl.load_workbook(original), openpyxl.load_workbook(enriched)
                    self.assertEqual(a.sheetnames, b.sheetnames)
                    for sheet in a:
                        other = b[sheet.title]
                        self.assertEqual(sheet.max_row, other.max_row)
                        if sheet.title != "resumen por tema": self.assertEqual(sheet.max_column, other.max_column)
                        if sheet.title == "detalle":
                            self.assertEqual(sorted(list(sheet.values)[1:], key=str), sorted(list(other.values)[1:], key=str))
                        for cells in sheet:
                            for cell in cells:
                                if sheet.title == "resumen general" and cell.column == 11: continue
                                if sheet.title != "detalle" or cell.row == 1:
                                    self.assertEqual(cell.value, other[cell.coordinate].value, (basis, sheet.title, cell.coordinate))
                                self.assertEqual(cell._style, other[cell.coordinate]._style)
                    headers = [c.value for c in b["resumen por tema"][1]]
                    self.assertEqual(headers[:10], [c.value for c in a["resumen por tema"][1]])
                    self.assertIn("Invitada", headers)
                    self.assertNotIn("Invitada %", headers)
                    self.assertEqual(headers[headers.index("Invitada") + 1], "Importe Invitada")
                    self.assertEqual(b["detalle"].max_row, 3)
                    a.close(); b.close()


if __name__ == "__main__":
    unittest.main()
